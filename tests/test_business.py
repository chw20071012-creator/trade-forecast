import copy
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from trade_agent.business import BusinessService
from trade_agent.store import Store


class BusinessTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.store = Store(self.temp.name)
        self.service = BusinessService(self.store)
        self.product = self.store.put("products", {"name": "Canvas bag", "business": "bags", "unit": "pcs", "moq": "100", "purchase_unit_price": "8", "purchase_currency": "USD", "purchase_cost_confirmed": True, "demo": True}, "p1")
        self.customer = self.store.put("customers", {"name": "Demo store", "country": "US", "demo": True}, "c1")
        self.rule = self.service.create_profit_rule({"business": "bags", "product_id": "p1", "proposed_min_margin": "0.20", "suggested_margin": "0.25", "rationale": "Tested proposal"})

    def costs(self):
        return [
            {"category": "purchase", "amount": "8", "currency": "USD", "basis": "per_unit", "confirmed": True},
            {"category": "packaging", "amount": "1", "currency": "USD", "basis": "per_unit", "confirmed": True},
            {"category": "testing", "amount": "100", "currency": "USD", "basis": "order", "confirmed": True},
            *[{"category": c, "amount": "0", "currency": "USD", "basis": "order", "confirmed": True} for c in ("freight", "payment", "service")],
        ]

    def request(self, **extra):
        return {"product_id": "p1", "quantity": "100", "currency": "USD", "terms": "EXW", "profit_rule_id": self.rule["id"], "cost_items": self.costs(), **extra}

    def ready_quote(self, **extra):
        self.service.confirm_profit_rule(self.rule["id"])
        return self.service.quote(self.request(**extra))

    def order(self):
        quote = self.ready_quote()
        return self.service.create_order({"customer_id": "c1", "quote_id": quote["id"], "seller": {"name": "Demo seller", "address": "Example"}, "buyer": {"name": "Demo store", "address": "Example"}, "lines": [{"product_id": "p1", "description": "Canvas bag", "quantity": "100", "unit": "pcs", "unit_price": quote["suggested_price"], "net_weight": "20", "gross_weight": "25", "packages": "2"}]})

    def test_true_gross_margin_and_confirmation(self):
        pending = self.service.quote(self.request())
        self.assertFalse(pending["usable"])
        self.assertIsNone(pending["floor_price"])
        self.assertIn("confirmed_profit_rule", pending["missing_costs"])
        quote = self.ready_quote()
        # Cost 8 + 1 + 100/100 = 10; 20% gross margin is 10/(1-.2).
        self.assertEqual(Decimal(quote["cost_per_unit"]), Decimal("10"))
        self.assertEqual(quote["floor_price"], "12.50")
        self.assertEqual(quote["suggested_price"], "13.34")
        actual_margin = (Decimal(quote["floor_price"]) - Decimal("10")) / Decimal(quote["floor_price"])
        self.assertEqual(actual_margin, Decimal("0.20"))

    def test_missing_costs_cannot_be_hidden_by_category_override(self):
        quote = self.ready_quote(cost_items=self.costs()[:1], required_cost_categories=["purchase"])
        self.assertFalse(quote["usable"])
        self.assertIn("cost:freight", quote["missing_costs"])
        self.assertIsNone(quote["suggested_price"])

    def test_quantity_tiers_fixed_cost_fx_and_small_order(self):
        product = {**self.product, "purchase_currency": "CNY", "price_tiers": [{"min_quantity": "1", "unit_price": "80"}, {"min_quantity": "100", "unit_price": "70"}]}
        self.store.put("products", product, "p1")
        quote = self.ready_quote(cost_items=self.costs()[1:], fx_rates={"CNY/USD": "0.1"})
        self.assertEqual(Decimal(quote["cost_per_unit"]), Decimal("9"))
        smaller = self.service.quote(self.request(quantity="10"))
        self.assertFalse(smaller["usable"])
        self.assertIn("supplier_small_order_confirmation", smaller["missing_costs"])
        confirmed_small = self.service.quote(self.request(quantity="10", supplier_small_order_confirmed=True))
        self.assertTrue(confirmed_small["usable"])
        self.assertEqual(Decimal(confirmed_small["cost_per_unit"]), Decimal("19"))

    def test_missing_fx_and_revenue_percentage_fee(self):
        costs = self.costs()
        costs[0]["currency"] = "CNY"
        bad = self.ready_quote(cost_items=costs)
        self.assertFalse(bad["usable"])
        self.assertIn("Missing exchange rate CNY/USD", bad["missing_costs"])
        good = self.service.quote(self.request(revenue_fee_rate="0.05"))
        self.assertEqual(good["floor_price"], "13.34")
        self.assertGreaterEqual((Decimal(good["floor_price"]) * Decimal("0.95") - Decimal(10)) / Decimal(good["floor_price"]), Decimal("0.20"))

    def test_reject_retail_and_mismatched_comparables(self):
        evidence = {"product_id": "p1", "price": "11", "currency": "USD", "unit": "pcs", "terms": "EXW", "date": date.today().isoformat(), "source": "https://example.invalid/quote", "price_type": "export_quote"}
        retail = {**evidence, "price_type": "retail", "price": "80"}
        wrong_unit = {**evidence, "unit": "kg", "price": "1"}
        quote = self.ready_quote(market_comparables=[evidence, retail, wrong_unit])
        self.assertEqual(len(quote["market_comparables"]), 1)
        self.assertEqual(len(quote["rejected_comparables"]), 2)
        self.assertEqual(quote["suggested_price"], quote["floor_price"])
        self.assertTrue(any("below" in warning for warning in quote["warnings"]))

    def test_confirmation_bound_to_profit_rule_content(self):
        rule = self.service.confirm_profit_rule(self.rule["id"])
        rule["min_margin"] = "0.05"
        self.store.put("profit_rules", rule, rule["id"])
        with self.assertRaisesRegex(ValueError, "changed after confirmation"):
            self.service.quote(self.request())

    def test_mail_claim_is_not_verified_money_and_progress_needs_review(self):
        order = self.order()
        update = self.service.propose_update({"order_id": order["id"], "stage": "receipts", "source": "email", "source_ref": "msg1", "patch": {"status": "verified", "verified_amount": order["total"], "bank_verified": True}})
        self.assertEqual(self.store.get("orders", order["id"])["stages"]["receipts"]["status"], "pending")
        self.assertEqual(update["patch"]["status"], "claimed")
        applied = self.service.confirm_update(update["id"])
        self.assertEqual(applied["order"]["stages"]["receipts"]["status"], "claimed")
        self.assertNotIn("verified_amount", applied["order"]["stages"]["receipts"])
        self.assertEqual(self.service.confirm_update(update["id"])["order"]["revision"], applied["order"]["revision"])

    def test_manual_bank_verification_and_stale_updates(self):
        order = self.order()
        with self.assertRaisesRegex(ValueError, "Verified receipt requires"):
            self.service.propose_update({"order_id": order["id"], "stage": "receipts", "patch": {"status": "verified", "verified_amount": order["total"]}})
        first = self.service.propose_update({"order_id": order["id"], "stage": "production", "patch": {"status": "in_progress", "completion_percent": "25"}})
        second = self.service.propose_update({"order_id": order["id"], "stage": "logistics", "patch": {"status": "booked"}})
        self.service.confirm_update(first["id"])
        with self.assertRaisesRegex(ValueError, "Order changed"):
            self.service.confirm_update(second["id"])
        receipt = self.service.propose_update({"order_id": order["id"], "stage": "receipts", "patch": {"status": "verified", "verified_amount": order["total"], "bank_verified": True, "verified_by": "owner", "evidence_reference": "bank-statement-row-1"}})
        self.assertEqual(self.service.confirm_update(receipt["id"])["order"]["stages"]["receipts"]["verified_amount"], order["total"])

    def test_only_allowed_stage_fields_and_no_scheduler(self):
        order = self.order()
        with self.assertRaisesRegex(ValueError, "unsupported stage fields"):
            self.service.propose_update({"order_id": order["id"], "stage": "production", "patch": {"currency": "GBP"}})
        result = self.service.followups({"order_id": order["id"]})
        self.assertFalse(result["scheduled"])
        self.assertEqual(result["messages_sent"], 0)
        self.assertTrue(any(a["stage"] == "receipts" for a in result["actions"]))

    def test_real_draft_documents_consistency_and_tampering(self):
        order = self.order()
        result = self.service.generate_documents(order["id"])
        self.assertEqual(len(result["files"]), 5)
        self.assertTrue(all(Path(path).is_file() and Path(path).stat().st_size > 1000 for path in result["files"]))
        invoice = next(p for p in result["files"] if p.endswith("draft_invoice.xlsx"))
        packing = next(p for p in result["files"] if p.endswith("draft_packing_list.xlsx"))
        checked = self.service.verify_documents({"order_id": order["id"], "invoice_path": invoice, "packing_list_path": packing})
        self.assertEqual(checked["status"], "consistent", checked)
        workbook = load_workbook(invoice)
        workbook["Invoice"]["I2"] = "USD"  # currency retained
        workbook["Invoice"]["H2"] = "9999.00"  # amount
        workbook.save(invoice)
        broken = self.service.verify_documents({"order_id": order["id"], "invoice_path": invoice, "packing_list_path": packing})
        self.assertFalse(broken["ok"])
        self.assertTrue(any(i["code"] == "amount_arithmetic" for i in broken["issues"]))

    def test_formula_is_never_evaluated_and_missing_weights_pending(self):
        order = self.order()
        result = self.service.generate_documents(order["id"])
        invoice = next(p for p in result["files"] if p.endswith("draft_invoice.xlsx"))
        packing = next(p for p in result["files"] if p.endswith("draft_packing_list.xlsx"))
        workbook = load_workbook(invoice)
        workbook["Invoice"]["H2"] = "=E2*G2"
        workbook.save(invoice)
        checked = self.service.verify_documents({"invoice_path": invoice, "packing_list_path": packing})
        self.assertEqual(checked["status"], "needs_review")
        self.assertFalse(checked["formulas_executed"])
        self.assertTrue(any(i["code"] == "formula_not_evaluated" for i in checked["pending"]))

    def test_row_alignment_not_position_and_duplicate_identity(self):
        invoice = [{"line_id": "A", "description": "A", "quantity": "2", "unit": "pcs", "unit_price": "3", "amount": "6", "currency": "USD"}, {"line_id": "B", "description": "B", "quantity": "1", "unit": "pcs", "unit_price": "4", "amount": "4", "currency": "USD"}]
        packing = [{"line_id": l["line_id"], "description": l["description"], "quantity": l["quantity"], "unit": "pcs", "net_weight": "1", "gross_weight": "2", "packages": "1"} for l in reversed(invoice)]
        self.assertTrue(self.service.verify_documents({"invoice": invoice, "packing_list": packing})["ok"])
        packing.append(copy.deepcopy(packing[0]))
        result = self.service.verify_documents({"invoice": invoice, "packing_list": packing})
        self.assertTrue(any(i["code"] == "ambiguous_duplicate_line" for i in result["issues"]))

    def test_generated_literal_text_cannot_be_a_spreadsheet_formula(self):
        order = self.order()
        order["lines"][0]["description"] = '=HYPERLINK("https://example.invalid","x")'
        self.store.put("orders", order, order["id"])
        generated = self.service.generate_documents(order["id"])
        invoice = next(p for p in generated["files"] if p.endswith("draft_invoice.xlsx"))
        workbook = load_workbook(invoice, data_only=False)
        self.assertEqual(workbook["Invoice"]["C2"].data_type, "s")
        self.assertEqual(workbook["Invoice"]["C2"].value, order["lines"][0]["description"])

    def test_quote_binding_prevents_order_unit_terms_and_product_changes(self):
        quote = self.ready_quote()
        base = {"customer_id": "c1", "quote_id": quote["id"], "currency": "USD", "terms": "EXW", "lines": [{"product_id": "p1", "quantity": "100", "unit": "pcs", "unit_price": "13.34"}]}
        for changes in ({"terms": "FOB"}, {"currency": "EUR"}, {"lines": [{"product_id": "p1", "quantity": "10", "unit": "pcs", "unit_price": "13.34"}]}, {"lines": [{"product_id": "p1", "quantity": "100", "unit": "kg", "unit_price": "13.34"}]}, {"lines": [{"product_id": "p1", "quantity": "100", "unit": "pcs", "unit_price": "12.49"}]}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                self.service.create_order({**base, **changes})

    def test_missing_moq_and_different_pricing_unit_require_review(self):
        product = {**self.product}
        product.pop("moq")
        self.store.put("products", product, "p1")
        quote = self.ready_quote()
        self.assertFalse(quote["usable"])
        self.assertIn("supplier_MOQ_or_quantity_confirmation", quote["missing_costs"])
        with self.assertRaisesRegex(ValueError, "pricing unit differs"):
            self.service.quote(self.request(unit="kg"))

    def test_followups_use_actual_sent_dates_and_suppress_unsubscribed(self):
        self.store.put("customers", {"name": "Unsubscribed", "do_not_contact": True, "next_contact_date": "2026-01-01"}, "c-blocked")
        self.store.put("mail_batches", {"items": [{"state": "sent", "sent_at": "2026-09-01T12:00:00Z", "snapshot": {"customer_id": "c1"}}]})
        result = self.service.followups({"as_of": "2026-09-10", "days": 7})
        self.assertFalse(any(a["customer_id"] == "c-blocked" for a in result["actions"]))
        action = next(a for a in result["actions"] if a["customer_id"] == "c1")
        self.assertEqual(action["type"], "customer_followup")
        self.assertEqual(action["due_date"], "2026-09-08")

    def test_party_metadata_and_missing_weight_are_not_silent_success(self):
        order = self.order()
        order["seller"]["address"] = ""
        order["lines"][0]["gross_weight"] = None
        self.store.put("orders", order, order["id"])
        result = self.service.generate_documents(order["id"])
        invoice = next(p for p in result["files"] if p.endswith("draft_invoice.xlsx"))
        packing = next(p for p in result["files"] if p.endswith("draft_packing_list.xlsx"))
        checked = self.service.verify_documents({"invoice_path": invoice, "packing_list_path": packing})
        self.assertEqual(checked["status"], "needs_review")
        self.assertTrue(any(i.get("field") == "seller_address" for i in checked["pending"]))
        self.assertTrue(any(i.get("field") == "gross_weight" for i in checked["pending"]))

    def test_order_only_verifies_latest_existing_files_and_detects_stale_revision(self):
        order = self.order()
        with self.assertRaisesRegex(ValueError, "generate drafts"):
            self.service.verify_documents({"order_id": order["id"]})
        generated = self.service.generate_documents(order["id"])
        checked = self.service.verify_documents({"order_id": order["id"]})
        self.assertTrue(checked["ok"], checked)
        self.assertEqual(checked["source_document_set"], generated["id"])
        proposal = self.service.propose_update({"order_id": order["id"], "stage": "production", "patch": {"status": "in_progress"}})
        self.service.confirm_update(proposal["id"])
        stale = self.service.verify_documents({"order_id": order["id"]})
        self.assertEqual(stale["status"], "needs_review")
        self.assertTrue(any(p["code"] == "stale_document_revision" for p in stale["pending"]))
        # Verification reads rather than recreates the original export.
        self.assertEqual(len(self.store.get("orders", order["id"])["document_sets"]), 1)

    def test_partial_bank_receipt_cannot_mark_full_order_paid(self):
        order = self.order()
        proposal = self.service.propose_update({"order_id": order["id"], "stage": "receipts", "patch": {"status": "verified", "verified_amount": "100", "bank_verified": True, "verified_by": "owner", "evidence_reference": "bank-line"}})
        self.assertEqual(proposal["patch"]["status"], "partially_verified")
        self.assertEqual(self.service.confirm_update(proposal["id"])["order"]["stages"]["receipts"]["status"], "partially_verified")


if __name__ == "__main__":
    unittest.main()
