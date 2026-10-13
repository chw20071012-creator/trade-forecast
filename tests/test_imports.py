import csv
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from trade_agent.imports import ImportService
from trade_agent.store import Store


class ImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.store = Store(self.root / "data")
        self.service = ImportService(self.store)

    def test_template_headers_source_mapping_and_original_retained(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "客户名单"
        sheet.append(["客户ID", "客户名称", "国家", "邮箱", "客户类型"])
        sheet.append(["c-import-1", "Small merchant", "US", "BUYER@example.invalid", "个体商"])
        products = workbook.create_sheet("茶叶产品")
        products.append(["产品ID", "产品名称", "茶类", "计价单位", "采购单价", "采购币种", "MOQ", "检测与认证"])
        products.append(["p-import-1", "Tea", "绿茶", "kg", "12.50", "CNY", "5", ""])
        source = self.root / "source.xlsx"
        workbook.save(source)
        result = self.service.import_file(source)
        self.assertEqual(result["imported_count"], 2)
        self.assertEqual(result["status"], "imported")
        customer = self.store.get("customers", "c-import-1")
        self.assertEqual(customer["email"], "buyer@example.invalid")
        self.assertEqual(customer["customer_type"], "个体商")
        self.assertEqual(customer["source"]["mapping"]["name"], "客户名称")
        self.assertEqual(customer["source"]["original_values"]["邮箱"], "BUYER@example.invalid")
        product = self.store.get("products", "p-import-1")
        self.assertEqual(product["business"], "tea")
        self.assertEqual(product["purchase_unit_price"], "12.50")
        self.assertFalse(product["purchase_cost_confirmed"])
        self.assertNotIn("tests_and_certifications", product)
        source.unlink()
        self.assertTrue(Path(result["original_path"]).exists())

    def test_formulas_not_used_as_costs_and_no_arbitrary_kind_writes(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "箱包产品"
        sheet.append(["产品名称", "采购单价", "采购币种", "计价单位"])
        sheet.append(["Bag", "=SUM(100,200)", "CNY", "pcs"])
        source = self.root / "formula.xlsx"
        workbook.save(source)
        result = self.service.import_file(source)
        self.assertEqual(result["status"], "needs_review")
        self.assertFalse(result["formulas_executed"])
        product = self.store.list("products")[0]
        self.assertNotIn("purchase_unit_price", product)
        self.assertEqual(product["source"]["original_values"]["采购单价"], "=SUM(100,200)")
        with self.assertRaisesRegex(ValueError, "Only customers/products"):
            self.service.import_file(source, "profit_rules")

    def test_csv_aliases_and_explicit_id_merge_preserves_existing_state(self):
        self.store.put("customers", {"name": "Old", "history": [{"status": "contacted"}], "whatsapp": "+12345678900"}, "c1")
        source = self.root / "clients.csv"
        with source.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["customer_id", "company_name", "email"])
            writer.writerow(["c1", "New", "owner@example.invalid"])
        result = self.service.import_file(source)
        self.assertEqual(result["imported_count"], 1)
        customer = self.store.get("customers", "c1")
        self.assertEqual(customer["name"], "New")
        self.assertEqual(customer["history"], [{"status": "contacted"}])
        self.assertEqual(customer["whatsapp"], "+12345678900")
        self.assertEqual(len(customer["import_history"]), 1)

    def test_docx_extraction_is_pending_observation_not_product_price(self):
        source = self.root / "catalog.docx"
        document = Document()
        document.add_paragraph("Canvas bag: USD 12. Supplier claim, not confirmed price.")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Tea"
        table.rows[0].cells[1].text = "USD 8"
        document.save(source)
        result = self.service.import_file(source, "products")
        self.assertEqual(result["status"], "pending_review")
        self.assertEqual(result["imported_count"], 0)
        self.assertEqual(self.store.list("products"), [])
        observation = self.store.list("import_observations")[0]
        self.assertTrue(observation["not_approved_cost_data"])
        self.assertIn("Canvas bag", observation["text"])
        self.assertIn("USD 8", observation["text"])

    def test_image_has_explicit_pending_ocr_status(self):
        source = self.root / "scan.png"
        source.write_bytes(b"image-placeholder-for-format-routing")
        result = self.service.import_file(source)
        self.assertEqual(result["status"], "pending_ocr")
        self.assertEqual(result["imported_count"], 0)
        self.assertTrue(Path(result["original_path"]).exists())
        self.assertEqual(result["observations"][0]["status"], "pending_ocr")

    def test_invalid_numbers_never_become_zero_price(self):
        source = self.root / "invalid.csv"
        source.write_text("product_name,business,purchase_price,currency,moq,unit\nBag,bags,NaN,USD,-2,pcs\n", encoding="utf-8")
        result = self.service.import_file(source)
        self.assertEqual(result["status"], "needs_review")
        product = self.store.list("products")[0]
        self.assertNotIn("purchase_unit_price", product)
        self.assertNotIn("moq", product)
        self.assertEqual(sum(e["code"] == "invalid_number" for e in result["errors"]), 2)

    def test_duplicate_alias_columns_rejected_before_import(self):
        source = self.root / "duplicate.csv"
        source.write_text("company_name,name,email\nA,B,x@example.invalid\n", encoding="utf-8")
        result = self.service.import_file(source, "customers")
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["imported_count"], 0)
        self.assertEqual(self.store.list("customers"), [])

    def test_real_text_pdf_extraction_retains_page_and_never_creates_prices(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        source = self.root / "catalog.pdf"
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=200)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 150 Td (Canvas bag USD 12) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(source)
        result = self.service.import_file(source)
        self.assertEqual(result["status"], "pending_review")
        self.assertIn("Canvas bag USD 12", result["observations"][0]["text"])
        self.assertEqual(result["observations"][0]["pages"], 1)
        self.assertEqual(self.store.list("products"), [])

    def test_scanned_and_encrypted_pdf_have_explicit_status(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        scanned = self.root / "scan.pdf"
        writer.write(scanned)
        self.assertEqual(self.service.import_file(scanned)["status"], "pending_ocr")
        writer.encrypt("test-password")
        locked = self.root / "locked.pdf"
        writer.write(locked)
        self.assertEqual(self.service.import_file(locked)["status"], "encrypted")

    def test_known_credentials_are_redacted_from_import_record_values(self):
        source = self.root / "customers.csv"
        source.write_text("name,email,api_key\nA,a@example.invalid,do-not-persist-this-credential\n", encoding="utf-8")
        self.service.import_file(source, "customers")
        customer = self.store.list("customers")[0]
        self.assertEqual(customer["source"]["original_values"]["api_key"], "[redacted]")
        self.assertTrue(customer["source"]["untrusted_content"])


if __name__ == "__main__":
    unittest.main()
