"""Auditable pricing, order progress and draft documents.

This module performs arithmetic and records user approvals. It does not create
market evidence, sign contracts, transfer money or schedule customer contact.
All monetary values persisted in records are decimal strings.
"""
from __future__ import annotations

import copy
import hashlib
import json
import re
import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_CEILING, ROUND_HALF_UP
from pathlib import Path


CENT = Decimal("0.01")
STAGES = ("contract", "production", "documents", "logistics", "receipts", "aftersales", "reorder")
COST_CATEGORIES = ("purchase", "packaging", "testing", "freight", "payment", "service", "other")
STAGE_FIELDS = {
    "contract": {"status", "document_reference", "signed_at", "agreed_terms", "notes"},
    "production": {"status", "completion_percent", "estimated_completion", "completed_at", "supplier", "notes"},
    "documents": {"status", "document_references", "checked_at", "notes"},
    "logistics": {"status", "carrier", "tracking_number", "bill_of_lading", "departure_date", "arrival_date", "delivery_date", "notes"},
    "receipts": {"status", "amount", "currency", "claimed_at", "verified_amount", "verified_at", "verified_by", "bank_verified", "evidence_reference", "notes"},
    "aftersales": {"status", "issue", "claim_amount", "credit_note", "notes"},
    "reorder": {"status", "next_contact_date", "estimated_quantity", "related_order_id", "notes"},
}
STAGE_STATUSES = {
    "contract": {"draft", "sent", "signed", "cancelled"},
    "production": {"not_started", "in_progress", "completed", "on_hold"},
    "documents": {"pending", "draft", "reviewed", "released"},
    "logistics": {"pending", "booked", "in_transit", "delivered", "exception"},
    "receipts": {"pending", "claimed", "partially_verified", "verified", "overdue", "refunded"},
    "aftersales": {"none", "open", "resolving", "resolved"},
    "reorder": {"monitoring", "interested", "quoted", "ordered", "closed"},
}


def _decimal(value, field, *, minimum=None, maximum=None):
    if isinstance(value, bool) or value is None:
        raise ValueError(f"{field} must be a finite decimal number")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise ValueError(f"{field} must be a finite decimal number") from None
    if not result.is_finite():
        raise ValueError(f"{field} must be a finite decimal number")
    if abs(result) > Decimal("1000000000000000") or result.as_tuple().exponent < -12:
        raise ValueError(f"{field} exceeds supported numeric precision/range")
    if minimum is not None and result < minimum:
        raise ValueError(f"{field} must be >= {minimum}")
    if maximum is not None and result > maximum:
        raise ValueError(f"{field} must be <= {maximum}")
    return result


def _number(value):
    return format(value, "f")


def _money(value, *, ceiling=False):
    return _number(value.quantize(CENT, rounding=ROUND_CEILING if ceiling else ROUND_HALF_UP))


def _currency(value):
    result = str(value or "").upper().strip()
    if not re.fullmatch(r"[A-Z]{3}", result):
        raise ValueError("currency must be a three-letter currency code")
    return result


def _utc():
    return datetime.now(timezone.utc).isoformat()


def _digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()


def _date(value):
    if not value:
        return None
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _business(value):
    aliases = {"bags": "bags", "bag": "bags", "luggage": "bags", "箱包": "bags", "tea": "tea", "茶叶": "tea"}
    try:
        return aliases[str(value).lower()]
    except KeyError:
        raise ValueError("business must be bags or tea") from None


class BusinessService:
    def __init__(self, store):
        self.store = store

    def create_profit_rule(self, request):
        """Create a proposed gross-margin rule; it is inactive until confirmed."""
        product = self.store.get("products", request["product_id"]) if request.get("product_id") else None
        business = _business(request.get("business") or (product or {}).get("business", "bags"))
        preferences = self.store.settings().get("profit_defaults", {})
        if not isinstance(preferences, dict):
            preferences = {}
        business_preferences = preferences.get(business, {})
        if not isinstance(business_preferences, dict):
            business_preferences = {}
        # A transparent editable starting proposal, not a market optimum.
        minimum = _decimal(request.get("proposed_min_margin", request.get("min_margin", business_preferences.get("min_margin", "0.12"))), "proposed_min_margin", minimum=Decimal(0), maximum=Decimal("0.95"))
        suggested = _decimal(request.get("suggested_margin", request.get("target_margin", business_preferences.get("target_margin", "0.20"))), "suggested_margin", minimum=minimum, maximum=Decimal("0.95"))
        evidence = request.get("evidence", [])
        if not isinstance(evidence, list):
            raise ValueError("evidence must be a list")
        rule = self.store.put("profit_rules", {
            "business": business, "product_id": request.get("product_id"),
            "min_margin": _number(minimum), "proposed_min_margin": _number(minimum),
            "suggested_margin": _number(suggested), "target_margin": _number(suggested),
            "status": "pending", "confirmed": False,
            "rationale": str(request.get("rationale") or "Editable starting proposal; validate full costs and comparable quotations before adoption."),
            "evidence": copy.deepcopy(evidence),
            "basis": "user_or_ai_proposal" if request.get("evidence") or request.get("rationale") else "configurable_starting_assumption",
            "limitations": ["Gross margin equals (revenue - all included costs) / revenue.", "These percentages are not validated market-optimal margins."],
            "demo": bool(request.get("demo", (product or {}).get("demo", self.store.settings().get("demo_mode", True)))),
        })
        self.store.event("profit_rule_proposed", {"rule_id": rule["id"]})
        return rule

    def confirm_profit_rule(self, rule_id):
        with self.store.lock:
            rule = self.store.get("profit_rules", rule_id)
            if rule.get("confirmed"):
                return rule
            _decimal(rule["min_margin"], "min_margin", minimum=Decimal(0), maximum=Decimal("0.95"))
            _decimal(rule["suggested_margin"], "suggested_margin", minimum=Decimal(rule["min_margin"]), maximum=Decimal("0.95"))
            rule.update(status="confirmed", confirmed=True, confirmed_at=_utc())
            rule["approved_digest"] = _digest({k: rule[k] for k in ("business", "product_id", "min_margin", "suggested_margin", "evidence")})
            result = self.store.put("profit_rules", rule, rule_id)
            self.store.event("profit_rule_confirmed", {"rule_id": rule_id})
            return result

    @staticmethod
    def _convert(amount, source, target, rates):
        if source == target:
            return amount
        key = f"{source}/{target}"
        reverse = f"{target}/{source}"
        if key in rates:
            return amount * _decimal(rates[key], key, minimum=Decimal("0.000000000001"))
        if reverse in rates:
            return amount / _decimal(rates[reverse], reverse, minimum=Decimal("0.000000000001"))
        # Each rate may alternatively be an explicit source-to-quote conversion.
        if source in rates:
            return amount * _decimal(rates[source], source, minimum=Decimal("0.000000000001"))
        raise ValueError(f"Missing exchange rate {source}/{target}")

    def quote(self, request):
        product = self.store.get("products", request["product_id"])
        quantity = _decimal(request.get("quantity"), "quantity", minimum=Decimal("0.000001"))
        currency = _currency(request.get("currency", "USD"))
        terms = str(request.get("terms") or "").strip().upper()
        unit = str(request.get("unit") or product.get("unit") or product.get("pricing_unit") or "").strip()
        product_unit = str(product.get("unit") or product.get("pricing_unit") or "").strip()
        if product_unit and unit != product_unit:
            raise ValueError("Quote pricing unit differs from supplier cost unit; convert supplier data explicitly first")
        rates = request.get("fx_rates", {})
        if not isinstance(rates, dict):
            raise ValueError("fx_rates must be an object")
        missing = []
        warnings = []
        if not terms:
            missing.append("trade_terms")
        if not unit:
            missing.append("pricing_unit")
        if product.get("quote_valid_until"):
            if not _date(product["quote_valid_until"]):
                missing.append("invalid_supplier_quote_expiry")
            elif _date(product["quote_valid_until"]) < date.today():
                missing.append("expired_supplier_quote")
        rule = self.store.get("profit_rules", request["profit_rule_id"]) if request.get("profit_rule_id") else None
        if not rule or not rule.get("confirmed"):
            missing.append("confirmed_profit_rule")
        if rule:
            if rule.get("product_id") and rule["product_id"] != product["id"]:
                raise ValueError("profit rule belongs to a different product")
            if rule.get("business") != _business(product.get("business", "bags")):
                raise ValueError("profit rule belongs to a different business")
            approved = _digest({k: rule[k] for k in ("business", "product_id", "min_margin", "suggested_margin", "evidence")})
            if rule.get("confirmed") and rule.get("approved_digest") != approved:
                raise ValueError("profit rule changed after confirmation; propose and confirm a new rule")
        minimum_margin = Decimal(rule["min_margin"]) if rule else Decimal("0.12")
        suggested_margin = Decimal(rule["suggested_margin"]) if rule else Decimal("0.20")
        items = copy.deepcopy(request.get("cost_items", []))
        if not isinstance(items, list):
            raise ValueError("cost_items must be a list")
        if not any(isinstance(i, dict) and i.get("category") == "purchase" for i in items):
            supplier_price = product.get("purchase_unit_price", product.get("purchase_price", product.get("cost")))
            tiers = product.get("price_tiers", [])
            eligible = [t for t in tiers if _decimal(t.get("min_quantity", 0), "min_quantity", minimum=Decimal(0)) <= quantity and (not t.get("max_quantity") or quantity <= _decimal(t["max_quantity"], "max_quantity", minimum=Decimal(0)))]
            if eligible:
                tier = max(eligible, key=lambda t: Decimal(str(t.get("min_quantity", 0))))
                supplier_price = tier.get("unit_price", tier.get("price"))
            if supplier_price is not None:
                items.insert(0, {"category": "purchase", "name": "Supplier purchase cost", "amount": supplier_price, "currency": product.get("purchase_currency", product.get("currency", currency)), "basis": "per_unit", "confirmed": bool(product.get("purchase_cost_confirmed", False)), "source": product.get("source", {}), "supplier_quote_date": product.get("quote_date")})
        required = request.get("required_cost_categories") or ["purchase", "packaging", "testing", "freight", "payment", "service"]
        if not isinstance(required, list) or any(category not in COST_CATEGORIES for category in required):
            raise ValueError("required_cost_categories contains unsupported categories")
        if "purchase" not in required:
            raise ValueError("purchase cost is always required")
        # Confirm zero where a category does not apply; silently omitting a
        # cost category cannot establish an actual profitable small order.
        required = sorted(set(required) | {"purchase", "packaging", "testing", "freight", "payment", "service"})
        included = []
        present = set()
        total_cost = Decimal(0)
        for i, item in enumerate(items):
            if not isinstance(item, dict):
                raise ValueError("each cost item must be an object")
            category = item.get("category", "other")
            if category not in COST_CATEGORIES:
                raise ValueError(f"unsupported cost category {category}")
            basis = item.get("basis", "per_unit")
            if basis not in ("per_unit", "order"):
                raise ValueError("cost basis must be per_unit or order")
            amount = _decimal(item.get("amount"), f"cost_items[{i}].amount", minimum=Decimal(0))
            item_currency = _currency(item.get("currency", currency))
            try:
                converted = self._convert(amount, item_currency, currency, rates)
            except ValueError as exc:
                missing.append(str(exc))
                included.append({**item, "amount": _number(amount), "currency": item_currency, "conversion_status": "missing_rate"})
                continue
            per_unit = converted if basis == "per_unit" else converted / quantity
            total_cost += per_unit
            present.add(category)
            if not item.get("confirmed"):
                missing.append(f"unconfirmed_cost:{category}:{i}")
            included.append({**item, "amount": _number(amount), "currency": item_currency, "basis": basis, "category": category, "quote_currency_per_unit": _number(per_unit)})
        missing.extend(f"cost:{category}" for category in required if category not in present)
        if product.get("moq") in (None, "") and not request.get("supplier_quantity_confirmed"):
            missing.append("supplier_MOQ_or_quantity_confirmation")
        moq = _decimal(product.get("moq", 0) or 0, "MOQ", minimum=Decimal(0))
        moq_unit = str(product.get("moq_unit") or product_unit).strip()
        if moq and moq_unit and moq_unit != unit:
            if request.get("moq_quantity_in_pricing_unit") is not None:
                moq = _decimal(request["moq_quantity_in_pricing_unit"], "moq_quantity_in_pricing_unit", minimum=Decimal(0))
            else:
                missing.append("MOQ_unit_conversion")
                moq = Decimal(0)
        small_order = quantity < moq
        if small_order and not request.get("supplier_small_order_confirmed"):
            missing.append("supplier_small_order_confirmation")
        if small_order:
            warnings.append("Quantity is below supplier MOQ; retain the opportunity and verify small-order costs and supply.")
        # A percentage payment fee is a fraction of sales, not a fixed purchase cost.
        fee_rate = _decimal(request.get("revenue_fee_rate", "0"), "revenue_fee_rate", minimum=Decimal(0), maximum=Decimal("0.95"))
        if minimum_margin + fee_rate >= 1 or suggested_margin + fee_rate >= 1:
            raise ValueError("margin plus revenue fee must be less than 1")
        provisional_floor = total_cost / (1 - minimum_margin - fee_rate)
        provisional_target = total_cost / (1 - suggested_margin - fee_rate)
        comparable_prices = []
        accepted = []
        rejected = []
        comparables = request.get("market_comparables", [])
        if not isinstance(comparables, list):
            raise ValueError("market_comparables must be a list")
        for index, comp in enumerate(comparables):
            reasons = []
            if not isinstance(comp, dict):
                rejected.append({"index": index, "reasons": ["not an object"]})
                continue
            if comp.get("price_type") not in ("export_quote", "transaction"):
                reasons.append("price type is not an export quote or transaction")
            if comp.get("product_id") != product["id"] and not comp.get("specifications_match_confirmed"):
                reasons.append("product specifications not confirmed comparable")
            if str(comp.get("unit", "")) != unit:
                reasons.append("different or unknown pricing unit")
            if str(comp.get("terms", "")).upper() != terms:
                reasons.append("different or unknown trade terms")
            if not comp.get("source") or not _date(comp.get("date")):
                reasons.append("missing source or valid date")
            elif _date(comp["date"]) > date.today():
                reasons.append("evidence date is in the future")
            elif (date.today() - _date(comp["date"])).days > int(request.get("comparable_max_age_days", 180)):
                reasons.append("evidence is too old")
            if comp.get("quantity_min") is not None and quantity < _decimal(comp["quantity_min"], "quantity_min", minimum=Decimal(0)):
                reasons.append("quantity below comparable tier")
            if comp.get("quantity_max") is not None and quantity > _decimal(comp["quantity_max"], "quantity_max", minimum=Decimal(0)):
                reasons.append("quantity above comparable tier")
            try:
                comparable_price = self._convert(_decimal(comp.get("price"), "comparable price", minimum=Decimal("0.000001")), _currency(comp.get("currency")), currency, rates)
            except ValueError as exc:
                reasons.append(str(exc))
            if reasons:
                rejected.append({"index": index, "reasons": reasons})
            else:
                comparable_prices.append(comparable_price)
                accepted.append({**comp, "quote_currency_price": _number(comparable_price)})
        if comparable_prices:
            sorted_prices = sorted(comparable_prices)
            middle = len(sorted_prices) // 2
            median = sorted_prices[middle] if len(sorted_prices) % 2 else (sorted_prices[middle - 1] + sorted_prices[middle]) / 2
            suggested_price = max(provisional_floor, median)
            if median < provisional_floor:
                warnings.append("Comparable market price is below the approved profit floor; renegotiate costs or terms before accepting.")
            price_basis = "median_of_matched_comparables_subject_to_profit_floor"
        else:
            suggested_price = provisional_target
            price_basis = "confirmed_cost_and_margin" if not missing else "incomplete_cost_estimate"
            warnings.append("No validated comparable quote: the target is a cost-based starting quote, not a proven market-optimal price.")
        complete = not missing
        result = {
            "product_id": product["id"], "product_name": product.get("name", product.get("description", "")),
            "business": _business(product.get("business", "bags")), "quantity": _number(quantity), "unit": unit,
            "currency": currency, "terms": terms, "profit_rule_id": (rule or {}).get("id"),
            "status": "ready_for_review" if complete else "needs_cost_review", "usable": complete,
            "cost_items": included, "cost_per_unit": _number(total_cost), "total_cost": _money(total_cost * quantity),
            "min_margin": _number(minimum_margin), "target_margin": _number(suggested_margin), "revenue_fee_rate": _number(fee_rate),
            "floor_price": _money(provisional_floor, ceiling=True) if complete else None,
            "minimum_unit_price": _money(provisional_floor, ceiling=True) if complete else None,
            "suggested_price": _money(suggested_price, ceiling=True) if complete else None,
            "suggested_unit_price": _money(suggested_price, ceiling=True) if complete else None,
            "negotiation_range": {"min": _money(provisional_floor, ceiling=True), "max": _money(suggested_price, ceiling=True)} if complete else None,
            "provisional_estimate": {"floor_price": _money(provisional_floor, ceiling=True), "target_price": _money(suggested_price, ceiling=True), "not_a_profit_guarantee": True},
            "price_basis": price_basis, "market_comparables": accepted, "rejected_comparables": rejected,
            "missing_costs": sorted(set(missing)), "warnings": warnings, "fx_rates": {str(k): str(v) for k, v in rates.items()},
            "demo": bool(product.get("demo") or (rule or {}).get("demo") or any(c.get("demo") for c in comparables if isinstance(c, dict))),
            "valid_until": request.get("valid_until"), "cost_snapshot_digest": _digest({"items": included, "quantity": _number(quantity), "fx": rates, "terms": terms}),
        }
        saved = self.store.put("quotes", result)
        self.store.event("quote_created", {"quote_id": saved["id"], "usable": complete})
        return saved

    def create_order(self, request):
        customer = self.store.get("customers", request["customer_id"]) if request.get("customer_id") else None
        quote = self.store.get("quotes", request["quote_id"]) if request.get("quote_id") else None
        if quote and not quote.get("usable"):
            raise ValueError("Cannot create an order from a quote with unverified costs")
        if quote and quote.get("valid_until") and _date(quote["valid_until"]) and _date(quote["valid_until"]) < date.today():
            raise ValueError("Quote has expired; generate a current quote")
        lines = copy.deepcopy(request.get("lines") or [])
        if not lines and quote:
            lines = [{"product_id": quote["product_id"], "description": quote["product_name"], "quantity": quote["quantity"], "unit": quote["unit"], "unit_price": quote["suggested_price"]}]
        if not lines:
            raise ValueError("order requires at least one line")
        if len(lines) > 1000:
            raise ValueError("Order exceeds 1000 line limit")
        if quote and (len(lines) != 1 or lines[0].get("product_id") != quote["product_id"]):
            raise ValueError("Order lines must match the quoted product; quote other products separately")
        currency = _currency(request.get("currency") or (quote or {}).get("currency", "USD"))
        if quote and currency != quote["currency"]:
            raise ValueError("Order currency must match quote currency")
        order_terms = str(request.get("terms") or (quote or {}).get("terms") or "").strip().upper()
        if quote and order_terms != quote["terms"]:
            raise ValueError("Trade terms changed; re-quote applicable logistics costs")
        normalized = []
        demo = bool((customer or {}).get("demo") or (quote or {}).get("demo") or request.get("demo", self.store.settings().get("demo_mode", True)))
        used_ids = set()
        for index, line in enumerate(lines):
            if not isinstance(line, dict):
                raise ValueError("order lines must be objects")
            quantity = _decimal(line.get("quantity"), f"lines[{index}].quantity", minimum=Decimal("0.000001"))
            price = _decimal(line.get("unit_price"), f"lines[{index}].unit_price", minimum=Decimal(0)).quantize(CENT, rounding=ROUND_HALF_UP)
            product = self.store.get("products", line["product_id"]) if line.get("product_id") else None
            demo = demo or bool((product or {}).get("demo"))
            line_id = str(line.get("line_id") or f"L{index + 1}")
            if line_id in used_ids:
                raise ValueError("order line_id values must be unique")
            used_ids.add(line_id)
            item = {"line_id": line_id, "product_id": line.get("product_id", ""), "description": str(line.get("description") or (product or {}).get("name") or "待补产品说明 / DESCRIPTION REQUIRED"), "specification": str(line.get("specification") or (product or {}).get("description") or ""), "quantity": _number(quantity), "unit": str(line.get("unit") or (product or {}).get("unit") or ""), "unit_price": _money(price), "amount": _money(quantity * price), "currency": currency}
            if not item["unit"]:
                raise ValueError("order line pricing unit is required")
            for field in ("net_weight", "gross_weight", "packages"):
                if line.get(field) not in (None, ""):
                    value = _decimal(line[field], field, minimum=Decimal(0))
                    if field == "packages" and value != value.to_integral_value():
                        raise ValueError("packages must be a whole number")
                    item[field] = _number(value)
                else:
                    item[field] = None
            if item["net_weight"] is not None and item["gross_weight"] is not None and Decimal(item["net_weight"]) > Decimal(item["gross_weight"]):
                raise ValueError("net_weight cannot exceed gross_weight")
            if quote and item["product_id"] == quote["product_id"]:
                if item["unit"] != quote["unit"]:
                    raise ValueError("Order pricing unit must match the quote")
                if quantity != Decimal(quote["quantity"]):
                    raise ValueError("Quantity changed from quote; re-quote quantity tiers and fixed costs")
                if price < Decimal(quote["floor_price"]):
                    raise ValueError("Order unit price is below confirmed quote profit floor")
            normalized.append(item)
        total = sum((Decimal(l["amount"]) for l in normalized), Decimal(0))
        milestones = {stage: {"status": {"contract": "draft", "production": "not_started", "documents": "pending", "logistics": "pending", "receipts": "pending", "aftersales": "none", "reorder": "monitoring"}[stage]} for stage in STAGES}
        buyer = copy.deepcopy(request.get("buyer") or {})
        if not buyer and customer:
            buyer = {"name": customer.get("name", ""), "address": customer.get("address", ""), "country": customer.get("country", ""), "email": customer.get("email", "")}
        seller = copy.deepcopy(request.get("seller") or self.store.settings().get("seller", {}))
        if not isinstance(seller, dict) or not isinstance(buyer, dict):
            raise ValueError("buyer and seller must be objects")
        order = self.store.put("orders", {"customer_id": request.get("customer_id"), "quote_id": request.get("quote_id"), "currency": currency, "terms": order_terms, "lines": normalized, "total": _money(total), "seller": seller, "buyer": buyer, "po_number": str(request.get("po_number") or ""), "payment_terms": str(request.get("payment_terms") or ""), "delivery_terms": str(request.get("delivery_terms") or ""), "stages": milestones, "revision": 1, "status": "draft", "pricing_basis": "confirmed_quote" if quote else "manual_order_requires_profit_review", "demo": demo, "document_sets": [], "history": []})
        self.store.event("order_created", {"order_id": order["id"], "demo": demo})
        return order

    def propose_update(self, request):
        order = self.store.get("orders", request["order_id"])
        stage = str(request.get("stage", ""))
        if stage not in STAGES:
            raise ValueError("unsupported order stage")
        source = str(request.get("source", "manual"))
        if source not in ("manual", "email", "external"):
            raise ValueError("source must be manual, email or external")
        patch = copy.deepcopy(request.get("patch", {}))
        if not isinstance(patch, dict) or not patch:
            raise ValueError("patch must be a non-empty object")
        unexpected = set(patch) - STAGE_FIELDS[stage]
        if unexpected:
            raise ValueError("unsupported stage fields: " + ", ".join(sorted(unexpected)))
        if "status" in patch and patch["status"] not in STAGE_STATUSES[stage]:
            raise ValueError("unsupported stage status")
        warnings = []
        for field in ("completion_percent", "amount", "verified_amount", "claim_amount", "estimated_quantity"):
            if field in patch:
                value = _decimal(patch[field], field, minimum=Decimal(0), maximum=Decimal(100) if field == "completion_percent" else None)
                patch[field] = _number(value)
        if "currency" in patch:
            patch["currency"] = _currency(patch["currency"])
            if patch["currency"] != order["currency"]:
                raise ValueError("Receipt currency mismatch; convert explicitly before updating")
        if stage == "receipts":
            verifying = patch.get("status") in ("verified", "partially_verified", "refunded") or patch.get("bank_verified") or "verified_amount" in patch
            if verifying and source != "manual":
                # Confirming an extraction confirms what a correspondent said.
                # It does not confirm settlement in the bank account.
                claimed = patch.get("amount", patch.get("verified_amount"))
                patch = {k: v for k, v in patch.items() if k not in ("verified_amount", "verified_at", "verified_by", "bank_verified")}
                patch["status"] = "claimed"
                if claimed is not None:
                    patch["amount"] = claimed
                warnings.append("Email/external payment statement recorded as a claim, not verified receipt.")
            elif verifying:
                if patch.get("bank_verified") is not True or not patch.get("evidence_reference") or not patch.get("verified_by") or "verified_amount" not in patch:
                    raise ValueError("Verified receipt requires manual bank_verified=true, verified_amount, verified_by and evidence_reference")
                if Decimal(patch["verified_amount"]) > Decimal(order["total"]):
                    warnings.append("Verified amount exceeds order total; reconcile overpayment separately.")
                if patch.get("status") != "refunded":
                    patch["status"] = "verified" if Decimal(patch["verified_amount"]) >= Decimal(order["total"]) else "partially_verified"
                    warnings.append("Verified amount is cumulative bank-verified settlement for this order, not a new amount to add automatically.")
                patch["currency"] = order["currency"]
                patch["verified_at"] = patch.get("verified_at") or _utc()
        proposal = {"order_id": order["id"], "stage": stage, "patch": patch, "source": source, "source_ref": request.get("source_ref"), "note": str(request.get("note") or ""), "status": "pending", "confirmed": False, "base_revision": order.get("revision", 1), "warnings": warnings, "demo": bool(order.get("demo"))}
        proposal["proposal_digest"] = _digest({k: proposal[k] for k in ("order_id", "stage", "patch", "base_revision", "source")})
        saved = self.store.put("order_updates", proposal)
        self.store.event("order_update_proposed", {"update_id": saved["id"], "order_id": order["id"]})
        return saved

    def confirm_update(self, update_id):
        with self.store.lock:
            proposal = self.store.get("order_updates", update_id)
            order = self.store.get("orders", proposal["order_id"])
            if proposal.get("status") == "applied":
                return {"update": proposal, "order": order}
            digest = _digest({k: proposal[k] for k in ("order_id", "stage", "patch", "base_revision", "source")})
            if digest != proposal.get("proposal_digest"):
                raise ValueError("Proposal changed; create a new proposal before approving")
            if order.get("revision", 1) != proposal["base_revision"]:
                raise ValueError("Order changed since proposal; review and propose again")
            order["stages"][proposal["stage"]].update(proposal["patch"])
            order["revision"] = order.get("revision", 1) + 1
            order["history"].append({"update_id": update_id, "stage": proposal["stage"], "source": proposal["source"], "source_ref": proposal.get("source_ref"), "confirmed_at": _utc(), "approved_digest": digest})
            if proposal["stage"] == "contract" and proposal["patch"].get("status") == "signed":
                order["status"] = "active"
            if proposal["stage"] == "contract" and proposal["patch"].get("status") == "cancelled":
                order["status"] = "cancelled"
            saved_order = self.store.put("orders", order, order["id"])
            proposal.update(status="applied", confirmed=True, confirmed_at=_utc(), applied_revision=order["revision"])
            saved_update = self.store.put("order_updates", proposal, update_id)
            self.store.event("order_update_confirmed", {"order_id": order["id"], "update_id": update_id})
            return {"update": saved_update, "order": saved_order}

    def generate_documents(self, order_id):
        """Export one immutable draft snapshot; never infer missing business fields."""
        from docx import Document
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        order = self.store.get("orders", order_id)
        snapshot = copy.deepcopy(order)
        identifier = uuid.uuid4().hex
        folder = Path(self.store.data_dir) / "documents" / hashlib.sha256(str(order_id).encode()).hexdigest()[:20] / identifier
        folder.mkdir(parents=True, exist_ok=True)
        metadata = {"status": "DRAFT / 待审核", "order_id": order_id, "order_revision": str(order["revision"]), "snapshot_digest": _digest(snapshot), "po_number": order.get("po_number", ""), "seller_name": order["seller"].get("name", ""), "seller_address": order["seller"].get("address", ""), "buyer_name": order["buyer"].get("name", ""), "buyer_address": order["buyer"].get("address", ""), "currency": order["currency"], "trade_terms": order.get("terms", ""), "total_amount": order["total"], "amount_in_words": _amount_words(Decimal(order["total"]), order["currency"]), "weight_unit": "kg", "demo": "true" if order.get("demo") else "false", "created_at": _utc()}
        missing_fields = [name for name in ("seller_name", "seller_address", "buyer_name", "buyer_address", "trade_terms") if not metadata[name]]
        for line in order["lines"]:
            missing_fields.extend(f"{line['line_id']}.{name}" for name in ("net_weight", "gross_weight", "packages") if line.get(name) is None)
        files = []
        for kind in ("invoice", "packing_list"):
            headers = ["line_id", "product_id", "description", "specification", "quantity", "unit"]
            headers += ["unit_price", "amount", "currency"] if kind == "invoice" else ["net_weight", "gross_weight", "packages", "currency"]
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Invoice" if kind == "invoice" else "Packing List"
            sheet.append(headers)
            for line in order["lines"]:
                for index, key in enumerate(headers, 1):
                    cell = sheet.cell(sheet.max_row + 1 if index == 1 else sheet.max_row, index)
                    value = line.get(key)
                    # Textual decimal values preserve precision and never execute formulas.
                    cell.value = "" if value is None else str(value)
                    cell.data_type = "s"
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="245B78")
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(55, max(16, max(len(str(c.value or "")) for c in column) + 2))
            sheet.freeze_panes = "A2"
            metasheet = workbook.create_sheet("Metadata")
            metasheet.append(["field", "value"])
            for key, value in metadata.items():
                metasheet.append([key, str(value) if value != "" else "待补 / REQUIRED"])
                metasheet.cell(metasheet.max_row, 2).data_type = "s"
            metasheet.append(["missing_fields", ", ".join(missing_fields)])
            metasheet.append(["notice", "Draft only. Complete missing fields; review before use. No signature or certification is implied."])
            summary = workbook.create_sheet("Summary")
            summary.append(["field", "value"])
            units = {l["unit"] for l in order["lines"]}
            if len(units) == 1:
                summary.append(["total_quantity", _number(sum((Decimal(l["quantity"]) for l in order["lines"]), Decimal(0)))])
            else:
                for unit in sorted(units):
                    summary.append(["total_quantity_" + unit, _number(sum((Decimal(l["quantity"]) for l in order["lines"] if l["unit"] == unit), Decimal(0)))])
            if kind == "invoice":
                summary.append(["total_amount", order["total"]])
            else:
                for field in ("net_weight", "gross_weight", "packages"):
                    known = all(l.get(field) is not None for l in order["lines"])
                    summary.append([f"total_{field}", _number(sum((Decimal(l[field]) for l in order["lines"]), Decimal(0))) if known else "待补 / REQUIRED"])
            xlsx = folder / f"draft_{kind}.xlsx"
            workbook.save(xlsx)
            files.append(str(xlsx.resolve()))
            document = Document()
            document.add_heading(("DEMO " if order.get("demo") else "") + "DRAFT " + ("COMMERCIAL INVOICE" if kind == "invoice" else "PACKING LIST"), 0)
            document.add_paragraph("待审核草稿 / Review required. Missing fields must be completed before use.")
            if order.get("demo"):
                document.add_paragraph("演示数据，不代表真实交易 / Demonstration only; not a real transaction.")
            for key in ("order_id", "order_revision", "po_number", "seller_name", "seller_address", "buyer_name", "buyer_address", "currency", "trade_terms"):
                document.add_paragraph(f"{key}: {metadata[key] or '待补 / REQUIRED'}")
            table = document.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            for index, key in enumerate(headers):
                table.rows[0].cells[index].text = key
            for line in order["lines"]:
                cells = table.add_row().cells
                for index, key in enumerate(headers):
                    cells[index].text = str(line[key]) if line.get(key) not in (None, "") else "待补 / REQUIRED"
            if kind == "invoice":
                document.add_paragraph(f"TOTAL {order['currency']}: {order['total']}")
                document.add_paragraph(metadata["amount_in_words"])
            document.add_paragraph(f"Snapshot: {metadata['snapshot_digest']}. No signature, payment execution or certification is implied.")
            docx = folder / f"draft_{kind}.docx"
            document.save(docx)
            files.append(str(docx.resolve()))
        document = Document()
        document.add_heading(("DEMO " if order.get("demo") else "") + "DRAFT SALES CONTRACT / 销售合同草稿", 0)
        document.add_paragraph("This is a review worksheet. Parties must complete and approve all commercial and legal terms before signing.")
        for key in ("order_id", "order_revision", "po_number", "seller_name", "seller_address", "buyer_name", "buyer_address", "currency", "trade_terms"):
            document.add_paragraph(f"{key}: {metadata[key] or '待补 / REQUIRED'}")
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for i, key in enumerate(("Line", "Description / Specification", "Quantity / Unit", "Unit price", "Amount")):
            table.rows[0].cells[i].text = key
        for line in order["lines"]:
            row = table.add_row().cells
            values = [line["line_id"], line["description"] + " " + line["specification"], line["quantity"] + " " + line["unit"], line["unit_price"], line["amount"]]
            for i, value in enumerate(values):
                row[i].text = value
        document.add_paragraph(f"Total: {order['currency']} {order['total']}")
        document.add_paragraph(f"Payment terms: {order.get('payment_terms') or '待补 / REQUIRED'}")
        document.add_paragraph(f"Delivery terms: {order.get('delivery_terms') or '待补 / REQUIRED'}")
        for term in ("Product quality and acceptance criteria", "Testing / documented certifications", "Inspection and claim procedure", "Applicable law and dispute resolution", "Seller signature", "Buyer signature"):
            document.add_paragraph(term + ": 待补 / REQUIRED")
        document.add_paragraph("Snapshot: " + metadata["snapshot_digest"])
        contract = folder / "draft_contract.docx"
        document.save(contract)
        files.append(str(contract.resolve()))
        export = {"id": identifier, "status": "draft", "order_revision": order["revision"], "snapshot_digest": metadata["snapshot_digest"], "created_at": _utc(), "missing_fields": missing_fields, "files": files}
        with self.store.lock:
            current = self.store.get("orders", order_id)
            current.setdefault("document_sets", []).append(export)
            self.store.put("orders", current, order_id)
        self.store.event("draft_documents_generated", {"order_id": order_id, "export_id": identifier, "revision": order["revision"]})
        return {"order_id": order_id, **export, "demo": bool(order.get("demo"))}

    def _read_document(self, path, kind):
        from openpyxl import load_workbook
        resolved = Path(path).resolve()
        if not resolved.is_relative_to(Path(self.store.data_dir).resolve()):
            raise ValueError("Document must be an imported/generated file inside the data directory")
        if resolved.suffix.lower() != ".xlsx":
            raise ValueError("Deterministic verification currently accepts XLSX or structured rows")
        _check_archive(resolved)
        workbook = load_workbook(resolved, read_only=True, data_only=False)
        preferred = "Invoice" if kind == "invoice" else "Packing List"
        sheet = workbook[preferred] if preferred in workbook.sheetnames else workbook.worksheets[0]
        iterator = iter(sheet.iter_rows())
        first = next(iterator, ())
        aliases = {"行号": "line_id", "产品ID": "product_id", "产品名称": "description", "品名": "description", "规格": "specification", "数量": "quantity", "单位": "unit", "单价": "unit_price", "金额": "amount", "币种": "currency", "净重": "net_weight", "毛重": "gross_weight", "箱数": "packages"}
        headers = [aliases.get(str(cell.value or "").strip(), str(cell.value or "").strip().lower().replace(" ", "_")) for cell in first]
        rows, formulas = [], []
        for number, row in enumerate(iterator, 2):
            if all(cell.value in (None, "") for cell in row):
                continue
            item = {}
            for header, cell in zip(headers, row):
                if not header:
                    continue
                if cell.data_type == "f":
                    formulas.append({"sheet": sheet.title, "row": number, "field": header})
                    item[header] = None
                else:
                    item[header] = cell.value
            rows.append(item)
        metadata = {}
        for title in ("Metadata", "Summary"):
            if title in workbook.sheetnames:
                for row in workbook[title].iter_rows(min_row=2):
                    if len(row) >= 2 and row[0].value:
                        if row[1].data_type == "f":
                            formulas.append({"sheet": title, "row": row[1].row, "field": str(row[0].value)})
                        else:
                            metadata[str(row[0].value)] = row[1].value
        workbook.close()
        return rows, metadata, formulas

    def verify_documents(self, request):
        """Compare keyed lines and recalculate arithmetic; no spreadsheet evaluation."""
        request = copy.deepcopy(request)
        order = self.store.get("orders", request["order_id"]) if request.get("order_id") else None
        source_document_set = None
        if order and not any(request.get(key) for key in ("invoice", "invoice_rows", "packing_list", "packing_rows", "invoice_path", "packing_list_path")):
            sets = order.get("document_sets", [])
            if not sets:
                raise ValueError("No generated documents exist for this order; generate drafts before verification")
            latest = sets[-1]
            paths = latest.get("files", [])
            for filename, key in (("draft_invoice.xlsx", "invoice_path"), ("draft_packing_list.xlsx", "packing_list_path")):
                match = next((p for p in paths if isinstance(p, str) and Path(p).name == filename), None)
                if not match:
                    raise ValueError("Latest document set is incomplete; review or generate complete drafts")
                request[key] = match
            source_document_set = latest.get("id")
        issues, pending, skipped = [], [], []
        invoice = request.get("invoice", request.get("invoice_rows"))
        packing = request.get("packing_list", request.get("packing_rows"))
        invoice_meta = copy.deepcopy(request.get("invoice_metadata", {}))
        packing_meta = copy.deepcopy(request.get("packing_metadata", {}))
        for name, path in (("invoice", request.get("invoice_path")), ("packing_list", request.get("packing_list_path"))):
            if path:
                rows, meta, formulas = self._read_document(path, "invoice" if name == "invoice" else "packing")
                if name == "invoice":
                    invoice, invoice_meta = rows, meta
                else:
                    packing, packing_meta = rows, meta
                for formula in formulas:
                    pending.append({"code": "formula_not_evaluated", "document": name, **formula})
        if not isinstance(invoice, list) or not isinstance(packing, list) or not invoice or not packing:
            raise ValueError("Both invoice and packing_list must contain rows or XLSX paths")
        for field in ("order_id", "po_number", "seller_name", "seller_address", "buyer_name", "buyer_address", "currency", "trade_terms", "weight_unit"):
            a, b = invoice_meta.get(field), packing_meta.get(field)
            if (a is None and b is None) or (field == "po_number" and (a in (None, "") or str(a).startswith("待补")) and (b in (None, "") or str(b).startswith("待补"))):
                skipped.append({"code": "metadata_not_provided", "field": field})
            elif a in (None, "") or b in (None, "") or str(a).startswith("待补") or str(b).startswith("待补"):
                pending.append({"code": "metadata_missing", "field": field})
            elif str(a).strip() != str(b).strip():
                issues.append({"code": "metadata_mismatch", "field": field, "invoice": str(a), "packing_list": str(b)})
        if order:
            expected_metadata = {"order_id": order["id"], "po_number": order.get("po_number"), "seller_name": order.get("seller", {}).get("name"), "seller_address": order.get("seller", {}).get("address"), "buyer_name": order.get("buyer", {}).get("name"), "buyer_address": order.get("buyer", {}).get("address"), "currency": order["currency"], "trade_terms": order.get("terms")}
            for field, expected in expected_metadata.items():
                supplied = invoice_meta.get(field)
                if supplied not in (None, "") and not str(supplied).startswith("待补") and expected not in (None, "") and str(supplied).strip() != str(expected).strip():
                    issues.append({"code": "order_metadata_mismatch", "field": field})
        invoice_index = _index_lines(invoice, "invoice", issues)
        packing_index = _index_lines(packing, "packing_list", issues)
        if set(invoice_index) != set(packing_index):
            issues.append({"code": "line_set_mismatch", "invoice_only": sorted(set(invoice_index) - set(packing_index)), "packing_only": sorted(set(packing_index) - set(invoice_index))})
        actual_total = Decimal(0)
        quantity_total = Decimal(0)
        amounts_complete = True
        quantities_complete = True
        currencies = set()
        for key, inv in invoice_index.items():
            pack = packing_index.get(key)
            for field in ("description", "specification", "unit", "product_id"):
                if pack and inv.get(field) not in (None, "") and pack.get(field) not in (None, "") and str(inv[field]).strip() != str(pack[field]).strip():
                    issues.append({"code": f"{field}_mismatch", "line": key})
            qty = _verify_number(inv.get("quantity"), "invoice", key, "quantity", issues, pending, positive=True)
            price = _verify_number(inv.get("unit_price"), "invoice", key, "unit_price", issues, pending)
            amount = _verify_number(inv.get("amount"), "invoice", key, "amount", issues, pending)
            if qty is not None:
                quantity_total += qty
            else:
                quantities_complete = False
            if amount is not None:
                actual_total += amount
            else:
                amounts_complete = False
            if qty is not None and price is not None and amount is not None:
                calculated = (qty * price).quantize(CENT, rounding=ROUND_HALF_UP)
                if abs(calculated - amount) >= CENT:
                    issues.append({"code": "amount_arithmetic", "line": key, "actual": _number(amount), "expected": _money(calculated)})
            if inv.get("currency"):
                currencies.add(str(inv["currency"]).upper().strip())
            elif invoice_meta.get("currency"):
                currencies.add(str(invoice_meta["currency"]).upper().strip())
            else:
                pending.append({"code": "missing_currency", "line": key})
            if pack:
                pack_qty = _verify_number(pack.get("quantity"), "packing_list", key, "quantity", issues, pending, positive=True)
                if qty is not None and pack_qty is not None and qty != pack_qty:
                    issues.append({"code": "quantity_mismatch", "line": key, "invoice": _number(qty), "packing_list": _number(pack_qty)})
                if pack.get("currency") and inv.get("currency") and str(pack["currency"]).upper() != str(inv["currency"]).upper():
                    issues.append({"code": "currency_mismatch", "line": key})
                net = _verify_number(pack.get("net_weight"), "packing_list", key, "net_weight", issues, pending)
                gross = _verify_number(pack.get("gross_weight"), "packing_list", key, "gross_weight", issues, pending)
                if net is not None and gross is not None and net > gross:
                    issues.append({"code": "net_exceeds_gross", "line": key})
                packages = _verify_number(pack.get("packages"), "packing_list", key, "packages", issues, pending)
                if packages is not None and packages != packages.to_integral_value():
                    issues.append({"code": "non_integer_packages", "line": key})
                if pack.get("unit_price") not in (None, ""):
                    pack_price = _verify_number(pack["unit_price"], "packing_list", key, "unit_price", issues, pending)
                    if pack_price is not None and price is not None and pack_price != price:
                        issues.append({"code": "unit_price_mismatch", "line": key})
                else:
                    skipped.append({"code": "packing_unit_price_absent", "line": key, "reason": "Packing lists need not contain unit prices."})
        if len(currencies) > 1:
            issues.append({"code": "mixed_currencies", "currencies": sorted(currencies)})
        if invoice_meta.get("currency") and currencies and str(invoice_meta["currency"]).upper() not in currencies:
            issues.append({"code": "invoice_metadata_currency_mismatch"})
        for field, actual in (("total_amount", actual_total), ("total_quantity", quantity_total)):
            if field == "total_quantity" and len({str(row.get("unit") or "") for row in invoice}) > 1:
                skipped.append({"code": "mixed_units_quantity_total_not_verifiable", "reason": "Quantity totals must be grouped by pricing unit."})
                continue
            if (field == "total_amount" and not amounts_complete) or (field == "total_quantity" and not quantities_complete):
                skipped.append({"code": field + "_not_verifiable", "reason": "One or more source values are missing or formulas were not evaluated."})
                continue
            if invoice_meta.get(field) not in (None, ""):
                supplied = _verify_number(invoice_meta[field], "invoice", "TOTAL", field, issues, pending)
                if supplied is not None and abs(supplied - actual) >= (CENT if field == "total_amount" else Decimal("0.000001")):
                    issues.append({"code": f"{field}_mismatch", "expected": _number(actual), "actual": _number(supplied)})
        for field in ("net_weight", "gross_weight", "packages", "quantity"):
            meta_key = "total_" + field
            if packing_meta.get(meta_key) not in (None, ""):
                known = [_verify_number(row.get(field), "packing_list", str(i), field, [], [], positive=False) for i, row in enumerate(packing)]
                stated = _verify_number(packing_meta[meta_key], "packing_list", "TOTAL", meta_key, issues, pending)
                if all(v is not None for v in known) and stated is not None and sum(known, Decimal(0)) != stated:
                    issues.append({"code": f"{meta_key}_mismatch", "expected": _number(sum(known, Decimal(0))), "actual": _number(stated)})
        words = request.get("amount_in_words", invoice_meta.get("amount_in_words"))
        if words and len(currencies) == 1 and amounts_complete:
            expected = _amount_words(actual_total, next(iter(currencies)))
            if _normalize_words(words) != _normalize_words(expected):
                issues.append({"code": "amount_in_words_mismatch", "expected": expected, "actual": str(words), "note": "Uses this program's deterministic English wording; review different valid conventions manually."})
        elif not words:
            skipped.append({"code": "amount_in_words_absent"})
        if order:
            order_index = _index_lines(order["lines"], "order", issues)
            if set(order_index) != set(invoice_index):
                issues.append({"code": "order_line_set_mismatch"})
            for key, expected in order_index.items():
                inv = invoice_index.get(key)
                if not inv:
                    continue
                for field in ("quantity", "unit_price", "amount"):
                    actual = _verify_number(inv.get(field), "invoice", key, field, [], [])
                    if actual is not None and actual != Decimal(expected[field]):
                        issues.append({"code": "order_" + field + "_mismatch", "line": key})
                for field in ("product_id", "description", "specification", "unit"):
                    if str(inv.get(field) or "") != str(expected.get(field) or ""):
                        issues.append({"code": "order_" + field + "_mismatch", "line": key})
            if currencies and currencies != {order["currency"]}:
                issues.append({"code": "order_currency_mismatch"})
            if invoice_meta.get("order_revision") not in (None, "") and str(invoice_meta["order_revision"]) != str(order["revision"]):
                pending.append({"code": "stale_document_revision", "document_revision": invoice_meta["order_revision"], "order_revision": order["revision"]})
        status = "inconsistent" if issues else "needs_review" if pending else "consistent"
        result = {"status": status, "ok": status == "consistent", "issues": issues, "pending": pending, "skipped": skipped, "calculated_invoice_total": _money(actual_total) if amounts_complete else None, "checked_at": _utc(), "formulas_executed": False, "order_id": request.get("order_id"), "source_document_set": source_document_set, "limitations": ["Consistency checks do not authenticate suppliers, signatures, certificates or bank settlement.", "Weight values are line totals in kg; quantities must share the stated unit."]}
        self.store.event("documents_verified", {"order_id": request.get("order_id"), "status": status, "issues": len(issues), "pending": len(pending)})
        return result

    def followups(self, request):
        """Return actions for the current invocation; never enqueue messages."""
        today = _date(request.get("as_of")) or date.today()
        actions = []
        orders = self.store.list("orders")
        customers = self.store.list("customers")
        latest_contacts = {}
        for batch in self.store.list("mail_batches"):
            for item in batch.get("items", []):
                contact_date = _date(item.get("sent_at")) if item.get("state") == "sent" else None
                customer_id = item.get("snapshot", {}).get("customer_id")
                if contact_date and customer_id:
                    latest_contacts[customer_id] = max(contact_date, latest_contacts.get(customer_id, contact_date))
        for message in self.store.list("whatsapp_messages"):
            contact_date = _date(message.get("manual_status_at")) if message.get("user_reported_sent") else None
            if contact_date and message.get("customer_id"):
                customer_id = message["customer_id"]
                latest_contacts[customer_id] = max(contact_date, latest_contacts.get(customer_id, contact_date))
        for order in orders:
            if request.get("order_id") and order["id"] != request["order_id"]:
                continue
            if request.get("customer_id") and order.get("customer_id") != request["customer_id"]:
                continue
            if order.get("status") == "cancelled":
                continue
            stages = order.get("stages", {})
            checks = (
                ("contract", stages.get("contract", {}).get("status") in ("draft", "sent"), "Review commercial terms and request contract confirmation."),
                ("production", stages.get("production", {}).get("status") in ("not_started", "in_progress", "on_hold"), "Ask the supplier for production progress and completion date."),
                ("documents", stages.get("documents", {}).get("status") in ("pending", "draft"), "Complete missing weights and parties, then verify invoice and packing list."),
                ("logistics", stages.get("logistics", {}).get("status") in ("pending", "booked", "in_transit", "exception"), "Check shipment milestone, tracking and delivery date."),
                ("receipts", stages.get("receipts", {}).get("status") != "verified", "Reconcile payment claim against bank evidence before marking received."),
                ("aftersales", stages.get("aftersales", {}).get("status") in ("open", "resolving"), "Review the customer's reported issue and propose resolution for approval."),
                ("reorder", stages.get("logistics", {}).get("status") == "delivered" and stages.get("reorder", {}).get("status") in ("monitoring", "interested", "quoted"), "Ask about sell-through and next replenishment quantity; prepare a current quote."),
            )
            for stage, needed, action in checks:
                if needed:
                    actions.append({"type": "order_followup", "order_id": order["id"], "customer_id": order.get("customer_id"), "stage": stage, "action": action, "due_date": today.isoformat(), "requires_confirmation": True, "demo": bool(order.get("demo"))})
        if not request.get("order_id"):
            for customer in customers:
                if request.get("customer_id") and customer["id"] != request["customer_id"]:
                    continue
                if customer.get("do_not_contact") or customer.get("contact_status") in ("unsubscribed", "hard_bounce", "do_not_contact") or customer.get("status") in ("do_not_contact", "unsubscribed"):
                    continue
                due = _date(customer.get("next_contact_date"))
                last_contact = _date(customer.get("last_contact_at") or customer.get("last_contact_date")) or latest_contacts.get(customer["id"])
                if not due and last_contact:
                    days = int(request.get("days", self.store.settings().get("followup_days", 7)))
                    if not 1 <= days <= 365:
                        raise ValueError("followup days must be between 1 and 365")
                    due = last_contact + timedelta(days=days)
                if not due and not last_contact:
                    actions.append({"type": "initial_outreach", "customer_id": customer["id"], "action": "Check buyer fit and public business evidence, then prepare individually tailored outreach for review.", "due_date": today.isoformat(), "requires_confirmation": True, "demo": bool(customer.get("demo"))})
                elif due and due <= today:
                    actions.append({"type": "customer_followup", "customer_id": customer["id"], "action": "Review last contact and any reply, generate an individually grounded draft, and request batch approval.", "due_date": due.isoformat(), "requires_confirmation": True, "demo": bool(customer.get("demo"))})
        return {"as_of": today.isoformat(), "actions": actions, "scheduled": False, "messages_sent": 0}


def _index_lines(rows, document, issues):
    index = {}
    for number, row in enumerate(rows, 1):
        if not isinstance(row, dict):
            issues.append({"code": "invalid_line", "document": document, "row": number})
            continue
        if row.get("line_id"):
            key = "line:" + str(row["line_id"])
        elif row.get("product_id"):
            key = "product:" + str(row["product_id"]) + "|" + str(row.get("specification") or "")
        elif row.get("description"):
            key = "description:" + str(row["description"]).strip() + "|" + str(row.get("specification") or "")
        else:
            issues.append({"code": "missing_line_identity", "document": document, "row": number})
            continue
        if key in index:
            issues.append({"code": "ambiguous_duplicate_line", "document": document, "line": key})
        else:
            index[key] = row
    return index


def _verify_number(value, document, line, field, issues, pending, positive=False):
    if value in (None, "") or str(value).startswith("待补"):
        pending.append({"code": "missing_value", "document": document, "line": line, "field": field})
        return None
    if isinstance(value, str) and value.lstrip().startswith("="):
        pending.append({"code": "formula_not_evaluated", "document": document, "line": line, "field": field})
        return None
    try:
        return _decimal(value, field, minimum=Decimal("0.000001") if positive else Decimal(0))
    except ValueError:
        issues.append({"code": "invalid_number", "document": document, "line": line, "field": field})
        return None


def _check_archive(path):
    from zipfile import ZipFile
    if path.stat().st_size > 25 * 1024 * 1024:
        raise ValueError("Document exceeds 25 MiB")
    with ZipFile(path) as archive:
        if len(archive.infolist()) > 10000 or sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError("Expanded document exceeds limits")


def _integer_words(number):
    ones = ("ZERO ONE TWO THREE FOUR FIVE SIX SEVEN EIGHT NINE TEN ELEVEN TWELVE THIRTEEN FOURTEEN FIFTEEN SIXTEEN SEVENTEEN EIGHTEEN NINETEEN").split()
    tens = ("ZERO TEN TWENTY THIRTY FORTY FIFTY SIXTY SEVENTY EIGHTY NINETY").split()
    if number < 20:
        return ones[number]
    if number < 100:
        return tens[number // 10] + (" " + ones[number % 10] if number % 10 else "")
    if number < 1000:
        return ones[number // 100] + " HUNDRED" + (" " + _integer_words(number % 100) if number % 100 else "")
    for scale, name in ((10**12, "TRILLION"), (10**9, "BILLION"), (10**6, "MILLION"), (1000, "THOUSAND")):
        if number >= scale:
            return _integer_words(number // scale) + " " + name + (" " + _integer_words(number % scale) if number % scale else "")


def _amount_words(amount, currency):
    amount = amount.quantize(CENT, rounding=ROUND_HALF_UP)
    whole = int(amount)
    cents = int((amount - whole) * 100)
    return f"{currency} {_integer_words(whole)} AND {cents:02d}/100 ONLY"


def _normalize_words(words):
    return re.sub(r"[^A-Z0-9/]", "", str(words).upper())
