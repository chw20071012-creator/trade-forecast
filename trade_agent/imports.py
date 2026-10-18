"""Safe structured imports with retained originals and traceable extraction.

Only customer/product columns are mapped into those record kinds. PDF/DOCX
extractions and images become pending observations, never automatic cost data.
"""
from __future__ import annotations

import copy
import csv
import hashlib
import io
import re
import shutil
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import BadZipFile, ZipFile


CUSTOMER_ALIASES = {
    "客户ID": "id", "客户名称": "name", "国家": "country", "城市": "city", "客户类型": "customer_type", "官网": "website",
    "采购联系人": "contact_name", "联系人岗位": "contact_role", "邮箱": "email", "业务电话": "phone", "WhatsApp号码": "whatsapp",
    "WhatsApp状态": "whatsapp_status", "匹配产品": "matched_products", "来源类型": "source_type", "来源链接": "source_url", "采集日期": "collected_at", "备注": "notes",
    "customer_id": "id", "company": "name", "company_name": "name", "customer_name": "name", "contact": "contact_name", "contact_person": "contact_name", "contact_title": "contact_role", "role": "contact_role", "url": "website", "telephone": "phone", "whatsapp_number": "whatsapp", "type": "customer_type", "source": "source_url", "date": "collected_at", "address": "address", "status": "status", "next_contact_date": "next_contact_date",
}
PRODUCT_ALIASES = {
    "产品ID": "id", "产品名称": "name", "品类": "category", "茶类": "category", "规格说明": "description", "材质": "material", "尺寸": "dimensions", "工艺与定制": "customization", "包装": "packaging", "计价单位": "unit", "采购单价": "purchase_unit_price", "采购币种": "purchase_currency", "MOQ": "moq", "MOQ单位": "moq_unit", "供应商": "supplier", "报价日期": "quote_date", "报价有效期": "quote_valid_until", "来源链接": "source_url", "备注": "notes", "等级": "grade", "产地": "origin", "检测与认证": "tests_and_certifications",
    "product_id": "id", "product_name": "name", "specification": "description", "specifications": "description", "pricing_unit": "unit", "purchase_price": "purchase_unit_price", "cost": "purchase_unit_price", "currency": "purchase_currency", "quote_expiry": "quote_valid_until", "source": "source_url", "business": "business",
}
CUSTOMER_FIELDS = set(CUSTOMER_ALIASES.values()) | {"name", "country", "city", "website", "email", "phone", "id", "notes"}
PRODUCT_FIELDS = set(PRODUCT_ALIASES.values()) | {"name", "id", "notes"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}
ALLOWED_EXTENSIONS = {".xlsx", ".csv", ".pdf", ".docx"} | IMAGE_EXTENSIONS


def _plain(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _normalize_header(value):
    return str(value or "").strip().replace("\ufeff", "")


def _alias(header, kind):
    mappings = CUSTOMER_ALIASES if kind == "customers" else PRODUCT_ALIASES
    canonical = header.lower().replace(" ", "_").replace("-", "_")
    fields = CUSTOMER_FIELDS if kind == "customers" else PRODUCT_FIELDS
    return mappings.get(header, mappings.get(canonical, canonical if canonical in fields else None))


def _archive_check(path):
    try:
        with ZipFile(path) as archive:
            if len(archive.infolist()) > 10000 or sum(e.file_size for e in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Expanded document exceeds 100 MiB/10000 entry limit")
    except BadZipFile:
        raise ValueError("Invalid Office document archive") from None


class ImportService:
    def __init__(self, store):
        self.store = store

    def import_file(self, path, kind=None):
        if kind not in (None, "customers", "products"):
            raise ValueError("Only customers/products may be imported as structured records")
        source = Path(path).resolve()
        if not source.is_file():
            raise ValueError("Import source is not a file")
        suffix = source.suffix.lower()
        if suffix not in ALLOWED_EXTENSIONS:
            raise ValueError("Supported formats: XLSX, CSV, text PDF, DOCX and images")
        if source.stat().st_size > 25 * 1024 * 1024:
            raise ValueError("Import exceeds 25 MiB")
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        import_id = uuid.uuid4().hex
        safe_name = re.sub(r"[^\w.\-]", "_", source.name)[:120] or f"upload{suffix}"
        folder = Path(self.store.data_dir) / "imports" / import_id
        folder.mkdir(parents=True, exist_ok=True)
        original = folder / safe_name
        shutil.copyfile(source, original)
        result = {"id": import_id, "import_id": import_id, "file_name": source.name, "sha256": digest, "requested_kind": kind, "original_path": str(original.resolve()), "status": "processing", "imported_count": 0, "records": [], "errors": [], "observations": [], "files": [str(original.resolve())], "formulas_executed": False}
        self.store.put("imports", result, import_id)
        try:
            if suffix == ".xlsx":
                self._xlsx(original, kind, result)
            elif suffix == ".csv":
                self._csv(original, kind, result)
            elif suffix == ".docx":
                self._docx(original, result)
            elif suffix == ".pdf":
                self._pdf(original, result)
            else:
                self._observe(result, "image", "", "pending_ocr", "Image retained. OCR requires an optional local tool or AI transcription; extracted values need review.")
                result["status"] = "pending_ocr"
        except (ValueError, OSError, BadZipFile) as exc:
            result["status"] = "failed"
            result["errors"].append({"code": "parse_failed", "message": str(exc)[:500]})
        if result["status"] == "processing":
            result["status"] = "needs_review" if result["errors"] else "imported" if result["imported_count"] else "no_records"
        saved = self.store.put("imports", result, import_id)
        self.store.event("file_imported", {"import_id": import_id, "status": saved["status"], "count": saved["imported_count"]})
        return saved

    def _detect_kind(self, headers, sheet, requested):
        if requested:
            return requested
        if sheet == "说明" or sheet.lower() in ("instructions", "readme", "metadata", "summary"):
            return None
        if "客户" in sheet or sheet.lower() in ("customers", "customer", "clients") or "客户名称" in headers:
            return "customers"
        if "产品" in sheet or sheet.lower() in ("products", "bags", "tea", "product") or "产品名称" in headers:
            return "products"
        if any(h.lower() in ("company_name", "customer_name", "contact_name", "website", "email") for h in headers):
            return "customers"
        if any(h.lower() in ("product_name", "purchase_unit_price", "purchase_price", "moq") for h in headers):
            return "products"
        return None

    def _xlsx(self, original, requested, result):
        from openpyxl import load_workbook
        _archive_check(original)
        workbook = load_workbook(original, read_only=True, data_only=False)
        try:
            for sheet in workbook:
                if sheet.title == "说明" or sheet.title.lower() in ("instructions", "readme"):
                    continue
                iterator = iter(sheet.iter_rows())
                headers = [_normalize_header(c.value) for c in next(iterator, ())]
                kind = self._detect_kind(headers, sheet.title, requested)
                if not kind:
                    result["errors"].append({"code": "unknown_sheet", "sheet": sheet.title, "message": "Choose customers/products or use the supplied template headers."})
                    continue
                mapped = [_alias(h, kind) for h in headers]
                duplicates = {h for h in mapped if h and mapped.count(h) > 1}
                if duplicates:
                    result["errors"].append({"code": "duplicate_columns", "sheet": sheet.title, "fields": sorted(duplicates)})
                    continue
                for row_number, cells in enumerate(iterator, 2):
                    if row_number > 20001:
                        raise ValueError("Maximum 20000 rows per sheet")
                    if all(c.value in (None, "") for c in cells):
                        continue
                    values = [_plain(c.value) for c in cells]
                    formulas = {i for i, c in enumerate(cells) if c.data_type == "f"}
                    self._row(kind, headers, values, sheet.title, row_number, formulas, result)
        finally:
            workbook.close()

    def _csv(self, original, requested, result):
        raw = original.read_bytes()
        encoding = None
        text = None
        for candidate in ("utf-8-sig", "gb18030", "utf-16"):
            try:
                text = raw.decode(candidate)
                encoding = candidate
                break
            except UnicodeDecodeError:
                pass
        if text is None:
            raise ValueError("CSV encoding not recognized; save as UTF-8 CSV")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        headers = [_normalize_header(h) for h in next(reader, [])]
        kind = self._detect_kind(headers, "CSV", requested)
        if not kind:
            raise ValueError("Cannot infer CSV record type; specify customers or products")
        mapped = [_alias(h, kind) for h in headers]
        duplicates = {h for h in mapped if h and mapped.count(h) > 1}
        if duplicates:
            raise ValueError("CSV has duplicate mapped columns: " + ", ".join(sorted(duplicates)))
        result["encoding"] = encoding
        for row_number, values in enumerate(reader, 2):
            if row_number > 20001:
                raise ValueError("Maximum 20000 CSV rows")
            if not any(v.strip() for v in values):
                continue
            formulas = {i for i, v in enumerate(values) if v.lstrip().startswith("=")}
            self._row(kind, headers, values, "CSV", row_number, formulas, result)

    def _row(self, kind, headers, values, sheet, row_number, formulas, result):
        record = {}
        mapping = {}
        source_values = {}
        for index, (header, value) in enumerate(zip(headers, values)):
            sensitive_header = header.lower().replace(" ", "_").replace("-", "_")
            source_values[header or f"column_{index + 1}"] = "[redacted]" if sensitive_header in {"password", "api_key", "access_token", "refresh_token", "client_secret", "private_key", "密码", "密钥"} else _plain(value)
            field = _alias(header, kind)
            if not field:
                continue
            mapping[field] = header
            if index in formulas:
                result["errors"].append({"code": "formula_not_evaluated", "sheet": sheet, "row": row_number, "field": field})
                continue
            if value not in (None, ""):
                record[field] = str(value).strip() if not isinstance(value, bool) else value
        if not record.get("name"):
            result["errors"].append({"code": "missing_name", "sheet": sheet, "row": row_number})
            return
        record_id = record.pop("id", None)
        if record_id and not re.fullmatch(r"[A-Za-z0-9_.\-]{1,100}", str(record_id)):
            result["errors"].append({"code": "invalid_id", "sheet": sheet, "row": row_number})
            return
        source = {"import_id": result["import_id"], "file_name": result["file_name"], "sha256": result["sha256"], "sheet": sheet, "row": row_number, "mapping": mapping, "original_values": source_values, "untrusted_content": True}
        record["source"] = source
        record["import_status"] = "needs_review"
        record["demo"] = False
        if kind == "products":
            business = record.get("business") or ("tea" if "茶" in sheet or sheet.lower() == "tea" else "bags" if "箱包" in sheet or sheet.lower() == "bags" else None)
            if business in ("箱包", "bag", "luggage"):
                business = "bags"
            if business in ("茶叶",):
                business = "tea"
            if business not in ("bags", "tea"):
                result["errors"].append({"code": "missing_business", "sheet": sheet, "row": row_number, "message": "Product business must be bags or tea."})
                record["business"] = "unclassified"
            else:
                record["business"] = business
            for field in ("purchase_unit_price", "moq"):
                if field in record:
                    try:
                        value = Decimal(str(record[field]).replace(",", ""))
                        if not value.is_finite() or value < 0:
                            raise InvalidOperation
                        record[field] = format(value, "f")
                    except (InvalidOperation, ValueError):
                        result["errors"].append({"code": "invalid_number", "sheet": sheet, "row": row_number, "field": field})
                        record.pop(field, None)
            if record.get("purchase_currency"):
                record["purchase_currency"] = str(record["purchase_currency"]).upper()
                if not re.fullmatch(r"[A-Z]{3}", record["purchase_currency"]):
                    result["errors"].append({"code": "invalid_currency", "sheet": sheet, "row": row_number})
                    record.pop("purchase_currency")
            record["purchase_cost_confirmed"] = False
        elif "email" in record:
            if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", record["email"]):
                record["email_validation"] = "invalid_format"
            else:
                record["email"] = record["email"].lower()
                record["email_validation"] = "format_only_unverified"
        if record_id:
            try:
                existing = self.store.get(kind, record_id)
            except KeyError:
                existing = None
            if existing:
                # Explicit IDs are user-provided import identifiers. Preserve
                # state/history not described by the permitted import columns.
                merged = copy.deepcopy(existing)
                merged.update(record)
                merged.setdefault("import_history", []).append(source)
                record = merged
        saved = self.store.put(kind, record, record_id)
        result["imported_count"] += 1
        result["records"].append({"kind": kind, "id": saved["id"], "sheet": sheet, "row": row_number})

    def _observe(self, result, source_type, text, status, notice, **extra):
        observation = self.store.put("import_observations", {"import_id": result["import_id"], "source_type": source_type, "text": text, "status": status, "notice": notice, "original_path": result["original_path"], "file_name": result["file_name"], "sha256": result["sha256"], "not_approved_cost_data": True, "untrusted_content": True, **extra})
        result["observations"].append({"id": observation["id"], "status": status, "source_type": source_type, "text": text, **extra})

    def _docx(self, original, result):
        from docx import Document
        _archive_check(original)
        document = Document(original)
        sections = [p.text for p in document.paragraphs if p.text.strip()]
        for number, table in enumerate(document.tables, 1):
            sections.append(f"[Table {number}]")
            sections.extend("\t".join(c.text for c in row.cells) for row in table.rows)
        text = "\n".join(sections)
        if len(text) > 1000000:
            raise ValueError("Extracted text exceeds one million characters")
        status = "pending_review" if text.strip() else "pending_ocr"
        self._observe(result, "docx", text, status, "Extracted text is untrusted source material. AI may propose structured values; a user must confirm them.")
        result["status"] = status

    def _pdf(self, original, result):
        try:
            from pypdf import PdfReader
        except ImportError:
            self._observe(result, "pdf", "", "parser_unavailable", "Install the bundled pypdf dependency to extract text; original retained.")
            result["status"] = "parser_unavailable"
            return
        try:
            reader = PdfReader(original)
            if reader.is_encrypted:
                self._observe(result, "pdf", "", "encrypted", "Encrypted PDF retained. Upload an unlocked copy for extraction.")
                result["status"] = "encrypted"
                return
            if len(reader.pages) > 500:
                raise ValueError("PDF exceeds 500 pages")
            chunks = []
            total_length = 0
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                chunks.append(f"[Page {number}]\n{text}")
                total_length += len(text)
                if total_length > 1000000:
                    raise ValueError("Extracted text exceeds one million characters")
            extracted = "\n".join(chunks)
            status = "pending_review" if total_length else "pending_ocr"
            self._observe(result, "pdf", extracted if total_length else "", status, "Text extraction does not verify prices/specifications. Scanned pages require optional OCR or AI transcription.", pages=len(reader.pages))
            result["status"] = status
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("PDF text extraction failed: " + type(exc).__name__) from None
